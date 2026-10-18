"""Pacchetto modificabile: sorgente HTML autonomo, dati di prodotti e codici (CSV/XLSX), font.
Uso: python3 scripts/package.py (dopo layout.py ed editable_build.py)"""
import csv
import json
import re
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
B = ROOT / 'build'
OUT = ROOT.parent / 'output' / 'definitivo' / 'modificabile'


def html_source():
    dst = OUT / 'sorgente-html'
    if dst.exists():
        shutil.rmtree(dst)
    for d in ('css', 'fonts', 'img', 'assets'):
        (dst / d).mkdir(parents=True)
    html = (B / 'catalogo.html').read_text()
    refs = set(re.findall(r'(?:src="|url\()(img/[^")]+)', html))
    for r in refs:
        shutil.copy(B / r, dst / r)
    for a in set(re.findall(r'\.\./\.\./assets/([^"]+)', html)):
        shutil.copy(ROOT.parent / 'assets' / a, dst / 'assets' / a)
    html = html.replace('../../assets/', 'assets/').replace('../scripts/v2.css', 'css/v2.css')
    (dst / 'catalogo.html').write_text(html)
    css = (ROOT / 'scripts' / 'v2.css').read_text()
    (dst / 'css' / 'v2.css').write_text(css)
    for f in (ROOT / 'fonts').glob('*.woff2'):
        shutil.copy(f, dst / 'fonts' / f.name)
    (dst / 'LEGGIMI.txt').write_text(
        'Sorgente dell\'impaginato: apri catalogo.html in Chrome o Edge e stampa in PDF (A4, margini nessuno, grafica di sfondo attiva).\n'
        'Testi e tabelle sono nel file HTML, lo stile in css/v2.css, le immagini in img/ e il marchio in assets/.\n'
        'Per la versione con abbondanza imposta in css/v2.css la variabile --b: 3mm e il formato pagina 216 × 303 mm.\n')
    return len(refs)


def data():
    rows = json.load(open(B / 'products.json'))
    cols = ['sezione', 'gruppo', 'famiglia', 'codice', 'variante', 'materiale', 'personalizzabile', 'pagina', 'immagine']
    d = OUT / 'dati'
    d.mkdir(exist_ok=True)
    with open(d / 'prodotti-e-codici.csv', 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=cols, delimiter=';'); w.writeheader(); w.writerows(rows)
    wb = Workbook(); ws = wb.active; ws.title = 'Prodotti e codici'
    ws.append([c.capitalize() for c in cols])
    for r in rows:
        ws.append([r[c] for c in cols])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='0A2A39')
    for i, wdt in enumerate([28, 30, 34, 16, 46, 30, 16, 8, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    wb.save(d / 'prodotti-e-codici.xlsx')
    return len(rows)




def readme():
    """Guida ai file consegnati (output/definitivo/LEGGIMI.md)."""
    ex = json.load(open(OUT.parent / 'export.json'))
    rep = json.load(open(B / 'layout_report.json'))
    parts = '\n'.join(f"- `{p['file']}` — pagine {p['pages']}, {p['mb']} MB" for p in ex['parts'])
    rows = json.load(open(B / 'products.json'))
    reps = sorted({r['famiglia'] for r in rows if r['immagine']})
    low = [f'{u} ({d} dpi)' for f, u, d in rep['low_dpi'] if d < 200]
    text = f"""# Waste Management Group · Catalogo prodotti

## File
- `Waste-Management-Group-Catalogo-digitale.pdf` — {ex['pages']} pagine A4, indice e collegamenti cliccabili, segnalibri. {ex['digital_mb']} MB.
- `Waste-Management-Group-Catalogo-stampa-A4-abbondanza-3mm.pdf` — per stampa e plastificazione: 216 × 303 mm con 3 mm di abbondanza per lato, TrimBox A4 210 × 297 mm. {ex['print_mb']} MB.
{parts}
- `modificabile/Waste-Management-Group-Catalogo-modificabile.pptx` — una diapositiva A4 per pagina: testi, immagini e tabelle sono oggetti separati e modificabili.
- `modificabile/font/` — Inter (Light, Regular, Medium, SemiBold): da installare prima di aprire il file PowerPoint.
- `modificabile/sorgente-html/` — impaginato originale (HTML + CSS + immagini), da cui si rigenerano i PDF.
- `modificabile/dati/prodotti-e-codici.xlsx` e `.csv` — {len(rows)} righe: sezione, famiglia, codice originale, variante, materiale, personalizzazione, pagina.

## Carattere
Helvetica Light non è disponibile in questo ambiente (è un carattere con licenza). È stato sostituito da Inter, un sans-serif dal disegno neutro e analogo, nei pesi Light (testi), Regular, Medium e SemiBold (titoli e codici). Se disponete della licenza Helvetica, nella versione PowerPoint basta sostituire il carattere.

## Immagini
- Le fotografie dei prodotti sono scontornate e presentate su un fondale uniforme con la stessa ombra.
- Le illustrazioni schematiche del documento di partenza sono state sostituite da render 3D fedeli a forma, materiale, colore e misure documentati.
- Le immagini che rappresentano solo la famiglia riportano la nota «Immagine rappresentativa».
- Le fotografie più deboli sono state migliorate con super-risoluzione (ricostruzione dei dettagli, non semplice ingrandimento) e non sono mai stampate sotto i 180 dpi effettivi.
- Famiglie con immagine rappresentativa: {len(reps)}.
- Immagini ancora al limite della risoluzione (da sostituire con foto originali del produttore quando disponibili): {', '.join(low) if low else 'nessuna'}.

## Da verificare prima della stampa
- Contatti: telefono 334 712 4042 e indirizzo Zona Industriale, C.da La Bruca, Scalea (CS) sono ripresi dal materiale fornito; il profilo Instagram @wastemanagementgroup è stato verificato.
- Prezzi, anno, quantità di confezione e minimi di personalizzazione non sono riportati.
"""
    (OUT.parent / 'LEGGIMI.md').write_text(text)


if __name__ == '__main__':
    OUT.mkdir(parents=True, exist_ok=True)
    print('immagini sorgente', html_source(), '· righe prodotto', data())
    readme()
