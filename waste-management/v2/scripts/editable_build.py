"""Versione modificabile del catalogo: PowerPoint A4 verticale, una diapositiva per pagina,
con testi, immagini e tabelle come oggetti separati (da build/editable.json).
Uso: python3 scripts/editable_build.py"""
import hashlib
import json
import shutil
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Mm, Pt

ROOT = Path(__file__).resolve().parent.parent
B = ROOT / 'build'
OUT = ROOT.parent / 'output' / 'definitivo' / 'modificabile'
PAGES = json.load(open(B / 'editable.json'))
FONT = {300: 'Inter Light', 400: 'Inter', 500: 'Inter Medium', 600: 'Inter SemiBold', 700: 'Inter SemiBold'}
ALIGN = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT, 'end': PP_ALIGN.RIGHT}
NO_STYLE = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'   # tabella senza stile né griglia
_media = {}


def rgb(h):
    return RGBColor.from_string(h.lstrip('#').upper())


def font_name(w):
    return FONT[min(FONT, key=lambda k: abs(k - (w or 300)))]


def media(it):
    """File dell'immagine; le icone rasterizzate identiche diventano un solo file."""
    if it.get('file'):
        p = B / 'editable_media' / it['file']
        h = hashlib.sha1(p.read_bytes()).hexdigest()
        return _media.setdefault(h, p)
    return ROOT.parent / it['src'].lstrip('/')


def add_img(slide, it):
    path = media(it)
    pic = slide.shapes.add_picture(str(path), Mm(it['x']), Mm(it['y']), Mm(it['w']), Mm(it['h']))
    if it.get('fit') == 'cover':      # ritaglio a riempimento, come nell'impaginato
        with Image.open(path) as im:
            nw, nh = im.size
        ar_box, ar = it['w'] / it['h'], nw / nh
        posy = 0.3 if '30%' in (it.get('pos') or '') else 0.5
        if ar > ar_box:
            c = 1 - ar_box / ar; pic.crop_left = pic.crop_right = c / 2
        else:
            c = 1 - ar / ar_box; pic.crop_top = c * posy; pic.crop_bottom = c * (1 - posy)
    return pic


def add_rect(slide, it):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if it['radius'] > 0.2 else MSO_SHAPE.RECTANGLE,
                                 Mm(it['x']), Mm(it['y']), Mm(it['w']), Mm(it['h']))
    if it['radius'] > 0.2:
        shp.adjustments[0] = min(0.5, it['radius'] / max(0.1, min(it['w'], it['h'])))
    shp.fill.solid(); shp.fill.fore_color.rgb = rgb(it['fill'])
    if it.get('alpha', 1) < 0.99:
        sf = shp.fill._xPr.find(qn('a:solidFill'))[0]
        a = etree.SubElement(sf, qn('a:alpha')); a.set('val', str(int(it['alpha'] * 100000)))
    shp.line.fill.background(); shp.shadow.inherit = False
    shp.text_frame.text = ''
    return shp


def style_run(r, s):
    f = r.font
    f.name = font_name(s.get('weight')); f.size = Pt(round(s['size'], 1)); f.color.rgb = rgb(s['color'])
    f.bold = False
    if s.get('ls'):
        r._r.get_or_add_rPr().set('spc', str(int(round(s['ls'] * 0.75 * 100))))


def add_text(slide, it):
    pad = 0.9
    w = it['w'] * 1.03 + pad
    x = it['x'] - (w - it['w']) * (1 if it['align'] in ('right', 'end') else 0.5 if it['align'] == 'center' else 0)
    tb = slide.shapes.add_textbox(Mm(x), Mm(it['y'] - 0.3), Mm(w), Mm(it['h'] + 0.6))
    tf = tb.text_frame
    tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE; tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras, cur = [], []
    for r in it['runs']:
        if r.get('br'):
            paras.append(cur); cur = []
        else:
            cur.append(r)
    paras.append(cur)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN.get(it['align'], PP_ALIGN.LEFT)
        p.line_spacing = round(max(0.8, it['lh'] / 1.21), 2)
        for s in runs:
            if not s.get('text'):
                continue
            r = p.add_run(); r.text = s['text']; style_run(r, s)
    url = next((s['link']['url'] for s in it['runs'] if s.get('link') and s['link'].get('url')), None)
    if url:   # collegamento sull'intera casella: il testo mantiene il suo colore
        tb.click_action.hyperlink.address = url
    return tb


def set_borders(cell, color, w_pt):
    """Solo il filetto inferiore; gli elementi lnL, lnR, lnT, lnB vanno in testa a tcPr, in quest'ordine."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('lnL', 'lnR', 'lnT', 'lnB'):
        for old in tcPr.findall(qn(f'a:{tag}')):
            tcPr.remove(old)
    for k, tag in enumerate(('lnL', 'lnR', 'lnT', 'lnB')):
        ln = etree.Element(qn(f'a:{tag}'), w=str(int((w_pt if tag == 'lnB' else 0) * 12700)))
        if tag == 'lnB':
            sf = etree.SubElement(ln, qn('a:solidFill')); c = etree.SubElement(sf, qn('a:srgbClr')); c.set('val', color.lstrip('#').upper())
        else:
            etree.SubElement(ln, qn('a:noFill'))
        tcPr.insert(k, ln)


def add_table(slide, it):
    rows, cols = len(it['rows']), max(len(r) for r in it['rows'])
    gf = slide.shapes.add_table(rows, cols, Mm(it['x']), Mm(it['y']), Mm(it['w']), Mm(it['h']))
    tbl = gf.table
    gf._element.graphic.graphicData.tbl.tblPr.find(qn('a:tableStyleId')).text = NO_STYLE
    tbl.first_row = False; tbl.horz_banding = False
    for j, wmm in enumerate(it['widths'][:cols]):
        tbl.columns[j].width = Mm(wmm)
    for i, hmm in enumerate(it['rowH']):
        tbl.rows[i].height = Mm(hmm)
    for i, row in enumerate(it['rows']):
        for j, c in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.background()
            cell.margin_left = 0; cell.margin_right = Mm(1.5); cell.margin_top = Mm(0.6 if i else 0); cell.margin_bottom = Mm(0.6)
            cell.vertical_anchor = MSO_ANCHOR.TOP
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; r = p.add_run(); r.text = c['text'].upper() if c.get('upper') else c['text']
            style_run(r, dict(size=c['size'], weight=c['weight'], color=c['color'], ls=c.get('ls')))
            set_borders(cell, it['line'] if i == 0 else it['rule'], 0.8 if i == 0 else 0.45)
    return gf


def main():
    if OUT.exists():
        shutil.rmtree(OUT)
    (OUT / 'font').mkdir(parents=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Mm(210), Mm(297)
    blank = prs.slide_layouts[6]
    links = []
    for pg in PAGES:
        s = prs.slides.add_slide(blank)
        s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(pg['bg'])
        for it in pg['items']:
            if it['t'] == 'rect':
                shp = add_rect(s, it)
            elif it['t'] == 'img':
                shp = add_img(s, it)
            elif it['t'] == 'table':
                shp = add_table(s, it)
            else:
                shp = add_text(s, it)
            lk = it.get('link') or next((r['link'] for r in it.get('runs', []) if r.get('link')), None)
            if lk and lk.get('page') and it['t'] != 'table':
                links.append((shp, lk['page']))
    for shp, n in links:
        if 1 <= n <= len(prs.slides):
            shp.click_action.target_slide = prs.slides[n - 1]
    prs.core_properties.title = 'Waste Management Group · Catalogo prodotti'
    prs.core_properties.author = 'Waste Management Group'
    path = OUT / 'Waste-Management-Group-Catalogo-modificabile.pptx'
    prs.save(path)
    for f in (ROOT / 'fonts' / 'static').glob('*.ttf'):
        shutil.copy(f, OUT / 'font' / f.name)
    print('pptx', path.stat().st_size // 1_000_000, 'MB', len(prs.slides), 'diapositive')


if __name__ == '__main__':
    main()
